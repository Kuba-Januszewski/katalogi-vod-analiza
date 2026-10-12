"""
Lekki runner SQL oparty na silniku DuckDB do analityki relacyjnej bazy VOD.

Projekt: Porównawcza analiza katalogów platform VOD (Amazon Prime, Disney+, HBO Max, Netflix)
Etap: Faza 4 - Analiza danych w SQL (DuckDB Runner)
"""

import sys
from pathlib import Path
import duckdb
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd

# Konfiguracja bezpiecznego kodowania konsoli dla środowiska Windows
if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

# Konfiguracja formatowania tabel Pandas w terminalu
pd.set_option("display.max_columns", None)
pd.set_option("display.width", 1000)

# Mapowanie bloków zapytań na arkusze Excela
BLOCK_SHEET_MAPPING = {
    "BLOK 1": "1_Formaty",
    "BLOK 2": "2_Jakosc",
    "BLOK 3": "3_Wylacznosc",
    "BLOK 4": "4_Gatunki_Wiek",
    "BLOK 5": "5_Tworcy",
}


def extract_queries_from_sql(sql_content: str) -> list[tuple[str, str, str]]:
    """
    Parsuje plik SQL na pojedyncze zapytania wraz z ich nagłówkami bloków i KPI.
    Zwraca listę krotek: (nagłówek_bloku, tytuł_kpi, treść_zapytania).
    """
    statements = []
    current_block = "BLOK ANALITYCZNY"
    raw_blocks = [block.strip() for block in sql_content.split(";") if block.strip()]

    for block in raw_blocks:
        lines = block.splitlines()
        kpi_title = ""
        sql_lines = []

        for line in lines:
            stripped = line.strip()
            if stripped.startswith("--"):
                clean_comment = stripped.lstrip("-").strip()
                if "BLOK" in clean_comment:
                    current_block = clean_comment
                elif "KPI" in clean_comment:
                    kpi_title = clean_comment
            elif stripped:
                sql_lines.append(line)

        query = "\n".join(sql_lines).strip()
        if query:
            if not kpi_title:
                kpi_title = "ZAPYTANIE ANALITYCZNE"
            statements.append((current_block, kpi_title, query))

    return statements


def print_formatted_table(df: pd.DataFrame) -> None:
    """Formatowanie i estetyczne wypisanie tabeli wynikowej w terminalu."""
    formatted_df = df.copy()
    for col in formatted_df.columns:
        if formatted_df[col].dtype == "float64":
            formatted_df[col] = formatted_df[col].apply(lambda x: f"{x:.2f}")

    table_str = formatted_df.to_string(index=False)
    print(table_str)


def export_reports_to_excel(
    executed_reports: list[tuple[int, str, str, pd.DataFrame]], 
    output_path: Path
) -> None:
    """
    Eksportuje zestaw raportów z Bloków 1–5 do jednego sformatowanego pliku Excel (.xlsx)
    z podziałem na 5 arkuszy tematycznych, czytelnymi nagłówkami sekcji i odstępami między tabelami.
    """
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Usunięcie domyślnego pustego arkusza

    # Grupowanie raportów według docelowych arkuszy
    sheets_data: dict[str, list[tuple[int, str, pd.DataFrame]]] = {}
    for idx, block_name, kpi_name, df in executed_reports:
        sheet_name = "Inne"
        for b_key, s_name in BLOCK_SHEET_MAPPING.items():
            if b_key in block_name:
                sheet_name = s_name
                break

        if sheet_name not in sheets_data:
            sheets_data[sheet_name] = []
        sheets_data[sheet_name].append((idx, kpi_name, df))

    # Definicje stylów dla arkusza (profesjonalny styl analityczny)
    font_title = Font(name="Calibri", size=11, bold=True, color="1F4E79")
    fill_title = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

    font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    fill_header = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")

    font_data = Font(name="Calibri", size=10)
    fill_zebra = PatternFill(start_color="F2F5F8", end_color="F2F5F8", fill_type="solid")

    thin_border_side = Side(border_style="thin", color="D3D3D3")
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    header_border_side = Side(border_style="thin", color="1A252F")
    header_border = Border(left=header_border_side, right=header_border_side, top=header_border_side, bottom=header_border_side)

    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    align_center = Alignment(horizontal="center", vertical="center")

    # Generowanie poszczególnych arkuszy
    for sheet_name, reports in sheets_data.items():
        ws = wb.create_sheet(title=sheet_name)
        ws.views.sheetView[0].showGridLines = True
        current_row = 1

        for r_idx, r_title, r_df in reports:
            num_cols = max(len(r_df.columns), 2)

            # 1. Nagłówek sekcji (tytuł raportu)
            ws.row_dimensions[current_row].height = 25
            for col_i in range(1, num_cols + 1):
                c = ws.cell(row=current_row, column=col_i)
                c.fill = fill_title
                c.border = cell_border
                if col_i == 1:
                    c.value = f"RAPORT {r_idx}: {r_title}"
                    c.font = font_title
                    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)

            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=num_cols)
            current_row += 1

            # 2. Wiersz nagłówków kolumn
            ws.row_dimensions[current_row].height = 22
            for col_idx, col_name in enumerate(r_df.columns, start=1):
                c = ws.cell(row=current_row, column=col_idx, value=str(col_name))
                c.font = font_header
                c.fill = fill_header
                c.border = header_border
                c.alignment = align_center

            current_row += 1

            # 3. Wiersze danych
            for row_i, row_data in enumerate(r_df.itertuples(index=False), start=1):
                ws.row_dimensions[current_row].height = 19
                is_even = (row_i % 2 == 0)
                for col_idx, val in enumerate(row_data, start=1):
                    c = ws.cell(row=current_row, column=col_idx)
                    c.font = font_data
                    c.border = cell_border
                    if is_even:
                        c.fill = fill_zebra

                    if pd.isna(val):
                        c.value = "-"
                        c.alignment = align_center
                    elif isinstance(val, (int,)):
                        c.value = int(val)
                        c.number_format = "#,##0"
                        c.alignment = align_right
                    elif isinstance(val, float):
                        c.value = float(val)
                        c.number_format = "0.00"
                        c.alignment = align_right
                    elif isinstance(val, bool):
                        c.value = "TAK" if val else "NIE"
                        c.alignment = align_center
                    else:
                        c.value = str(val)
                        c.alignment = align_left
                current_row += 1

            # Odstęp 2 wierszy między tabelami w arkuszu
            current_row += 2

        # 4. Automatyczne dopasowanie szerokości kolumn (z wyłączeniem wiersza tytułowego)
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.value and cell.font != font_title:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 14), 45)

    wb.save(output_path)


def main() -> None:
    base_dir = Path(__file__).resolve().parent.parent
    titles_path = base_dir / "data" / "processed" / "vod_titles_cleaned.csv"
    credits_path = base_dir / "data" / "processed" / "vod_credits_cleaned.csv"
    sql_path = base_dir / "sql" / "queries.sql"
    excel_path = base_dir / "reports" / "vod_analysis_report.xlsx"

    print("=" * 80)
    print("DUCKDB SQL ANALYTICS RUNNER: ANALIZA BENCHMARKINGOWA RYNKU VOD")
    print("=" * 80)

    # Weryfikacja dostępności plików
    if not titles_path.exists():
        raise FileNotFoundError(f"Brak pliku danych: {titles_path}. Uruchom najpierw scripts/data_cleaning.py")
    if not credits_path.exists():
        raise FileNotFoundError(f"Brak pliku danych: {credits_path}. Uruchom najpierw scripts/data_cleaning.py")
    if not sql_path.exists():
        raise FileNotFoundError(f"Brak pliku SQL: {sql_path}")

    # Inicjalizacja połączenia DuckDB w pamięci
    con = duckdb.connect(database=":memory:")

    # Rejestracja widoków analitycznych na podstawie wyczyszczonych plików CSV
    titles_posix_path = titles_path.as_posix()
    credits_posix_path = credits_path.as_posix()

    con.execute(f"""
        CREATE OR REPLACE VIEW vod_titles AS 
        SELECT * FROM read_csv_auto('{titles_posix_path}', header=True);
    """)
    con.execute(f"""
        CREATE OR REPLACE VIEW vod_credits AS 
        SELECT * FROM read_csv_auto('{credits_posix_path}', header=True);
    """)

    # Walidacja wczytania widoków
    total_titles = con.execute("SELECT COUNT(*) FROM vod_titles;").fetchone()[0]
    total_credits = con.execute("SELECT COUNT(*) FROM vod_credits;").fetchone()[0]
    print(f"\n[OK] Zarejestrowano widok 'vod_titles' w DuckDB (liczba rekordów: {total_titles:,})")
    print(f"[OK] Zarejestrowano widok 'vod_credits' w DuckDB (liczba rekordów: {total_credits:,})")
    print(f"[OK] Wczytano plik zapytań: {sql_path.relative_to(base_dir)}\n")

    # Wczytanie i parsowanie zapytań z pliku .sql
    with open(sql_path, "r", encoding="utf-8") as f:
        sql_content = f.read()

    queries = extract_queries_from_sql(sql_content)

    # Wykonanie poszczególnych zapytań, prezentacja w konsoli i zebranie do eksportu
    executed_reports: list[tuple[int, str, str, pd.DataFrame]] = []
    last_block = ""

    for idx, (block_name, kpi_name, query) in enumerate(queries, start=1):
        if block_name != last_block:
            print("\n" + "#" * 80)
            print(f"## {block_name}")
            print("#" * 80 + "\n")
            last_block = block_name

        print("-" * 80)
        print(f"RAPORT {idx}: {kpi_name}")
        print("-" * 80)

        result_df = con.execute(query).df()
        print_formatted_table(result_df)
        print()

        executed_reports.append((idx, block_name, kpi_name, result_df))

    # Eksport zestawu raportów do pliku Excel
    print("=" * 80)
    print("EKSPORT RAPORTÓW DO PLIKU EXCEL (.XLSX)")
    print("=" * 80)
    export_reports_to_excel(executed_reports, excel_path)
    print(f"[OK] Pomyślnie wyeksportowano raporty do pliku:")
    print(f"     {excel_path.relative_to(base_dir)}")
    print(f"     Arkusze: {', '.join(BLOCK_SHEET_MAPPING.values())}\n")

    print("=" * 80)
    print("WSZYSTKIE ZAPYTANIA (BLOKI 1–5) ORAZ EKSPORT ZOSTAŁY WYKONANE POMYŚLNIE")
    print("=" * 80 + "\n")


if __name__ == "__main__":
    main()

